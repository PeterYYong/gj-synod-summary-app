import os
import uuid
import hashlib
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import streamlit as st
from openai import OpenAI
from docx import Document

APP_TITLE = "나눔 회의록 요약 챗봇"
DEFAULT_MODEL = "gpt-5.5"
PROMPT_DIR = Path("prompts")
PROMPT_FILES = {
    "나눔 1": PROMPT_DIR / "nanum1_prompt.txt",
    "나눔 2": PROMPT_DIR / "nanum2_prompt.txt",
}
MODE_SLUG = {"나눔 1": "nanum1", "나눔 2": "nanum2"}

# 사용자에게 노출하지 않는 고정 API 설정
REASONING_EFFORT = "medium"
REASONING_SUMMARY = "auto"
TEXT_VERBOSITY = "medium"
STORE_RESPONSES = False
USE_STREAMING = True

COMMON_STYLE_RULES = """
공통 문체 보정 규칙:
1. 선택된 나눔 프롬프트의 구조와 요구사항을 우선 따르십시오.
2. 발표자가 실제로 읽어 줄 수 있는 자연스러운 한국어 문장으로 작성하십시오.
3. 발화자의 실제 단어와 중요한 문장 맥락은 유지하되, 반복되는 내용은 같은 의미끼리 묶어 정리하십시오.
4. 본문에는 발화자 이름, 번호, ID, 라벨을 직접 넣지 마십시오.
5. 원문에 없는 맥락, 결론, 감정, 담당자, 기한을 만들지 마십시오.
6. 같은 종결 표현을 연속해서 반복하지 마십시오.
7. “소개되었습니다”, “언급되었습니다”, “이야기가 나왔습니다”, “말이 있었습니다”는 필요할 때만 쓰고, 한 문단 안에서 반복하지 마십시오.
8. “반응이 있었습니다”, “반응으로 이어졌습니다”, “울림으로 이어졌습니다”, “공감이 형성되었습니다”는 사용하지 마십시오.
9. “그 마음에 함께 머물러 주셨습니다”, “그 어려움에 함께 마음을 모아 주셨습니다” 같은 인위적 표현은 사용하지 마십시오.
10. 부정적 경험은 민원이나 비난처럼 쓰지 말고, 함께 살펴야 할 과제로 정리하십시오.
11. 기존 프롬프트에 참조표가 요구되어 있으면 참조표를 유지하십시오.

문체 재점검:
- 초안 작성 후, 문장이 너무 기계적으로 반복되는지 확인하십시오.
- “소개되었습니다/언급되었습니다/나왔습니다”가 반복되면 자연스러운 문장으로 다시 풀어 쓰십시오.
- 발표자가 소리 내어 읽었을 때 어색한 문장은 짧고 쉬운 문장으로 고치십시오.
""".strip()


def init_session_state() -> None:
    defaults = {
        "session_id": str(uuid.uuid4()),
        "authenticated": False,
        "selected_mode": "나눔 1",
        "raw_input_text": "",
        "combined_source_text": "",
        "current_summary": "",
        "chat_history": [],
        "prompt_version": "",
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "input_version": 0,
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def reset_current_session(keep_auth: bool = True) -> None:
    """작업 데이터만 초기화합니다. 이미 생성된 위젯 key는 직접 수정하지 않습니다."""
    authenticated = st.session_state.get("authenticated", False)
    selected_mode = st.session_state.get("selected_mode", "나눔 1")

    st.session_state["session_id"] = str(uuid.uuid4())
    st.session_state["selected_mode"] = selected_mode
    st.session_state["raw_input_text"] = ""
    st.session_state["combined_source_text"] = ""
    st.session_state["current_summary"] = ""
    st.session_state["chat_history"] = []
    st.session_state["prompt_version"] = ""
    st.session_state["created_at"] = datetime.now().isoformat(timespec="seconds")
    # text_area를 확실히 비우기 위해 새 widget key를 사용하게 합니다.
    st.session_state["input_version"] = int(st.session_state.get("input_version", 0)) + 1
    st.session_state["authenticated"] = authenticated if keep_auth else False


def get_secret_value(name: str, default: Optional[str] = None) -> Optional[str]:
    try:
        value = st.secrets.get(name)
        if value is not None:
            return str(value)
    except Exception:
        pass
    return os.environ.get(name, default)


def get_model_name() -> str:
    return get_secret_value("OPENAI_MODEL", DEFAULT_MODEL) or DEFAULT_MODEL


def get_openai_client() -> OpenAI:
    api_key = get_secret_value("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY가 설정되어 있지 않습니다. Streamlit Secrets 또는 환경변수에 설정하세요.")
    return OpenAI(api_key=api_key)


def check_password() -> bool:
    expected = get_secret_value("APP_PASSWORD_4DIGIT")
    if not expected:
        st.error("APP_PASSWORD_4DIGIT가 설정되어 있지 않습니다. 관리자에게 문의하세요.")
        return False

    st.title(APP_TITLE)
    st.caption("접속하려면 관리자에게 받은 4자리 비밀번호를 입력하세요.")

    with st.form("password_form"):
        password = st.text_input("4자리 비밀번호", type="password", max_chars=4)
        submitted = st.form_submit_button("접속")

    if submitted:
        if password == expected:
            st.session_state["authenticated"] = True
            st.success("인증되었습니다.")
            st.rerun()
        else:
            st.error("비밀번호가 올바르지 않습니다.")
    return False


@st.cache_data(show_spinner=False)
def load_prompt_from_file(path_str: str) -> str:
    path = Path(path_str)
    if not path.exists():
        raise FileNotFoundError(f"프롬프트 파일을 찾을 수 없습니다: {path}")
    text = path.read_text(encoding="utf-8").strip()
    if not text:
        raise ValueError(f"프롬프트 파일이 비어 있습니다: {path}")
    return text


def prompt_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:10]


def load_prompt(mode: str) -> Tuple[str, str, str]:
    path = PROMPT_FILES[mode]
    text = load_prompt_from_file(str(path))
    version = f"{path.name}:{prompt_hash(text)}"
    return text, str(path), version


def build_instructions(mode: str, extra: str = "") -> Tuple[str, str]:
    prompt_text, prompt_name, version = load_prompt(mode)
    st.session_state["prompt_version"] = version
    parts = [
        prompt_text,
        "\n\n────────────────────────\n공통 보정 규칙\n────────────────────────\n",
        COMMON_STYLE_RULES,
    ]
    if extra.strip():
        parts.extend(["\n\n────────────────────────\n추가 편집 지시\n────────────────────────\n", extra.strip()])
    return "\n".join(parts), prompt_name


def response_to_text(response) -> str:
    if hasattr(response, "output_text") and response.output_text:
        return response.output_text
    chunks: List[str] = []
    try:
        for item in getattr(response, "output", []) or []:
            for content in getattr(item, "content", []) or []:
                text = getattr(content, "text", None)
                if text:
                    chunks.append(text)
    except Exception:
        pass
    return "\n".join(chunks).strip()


def extract_stream_delta(event) -> str:
    event_type = getattr(event, "type", "")
    if event_type == "response.output_text.delta":
        return getattr(event, "delta", "") or ""
    delta = getattr(event, "delta", None)
    if isinstance(delta, str):
        return delta
    data = getattr(event, "data", None)
    if isinstance(data, dict):
        return str(data.get("delta") or data.get("text") or "")
    if data is not None:
        text = getattr(data, "delta", None) or getattr(data, "text", None)
        if text:
            return str(text)
    return ""


def build_response_kwargs(instructions: str, user_input: str) -> dict:
    return {
        "model": get_model_name(),
        "instructions": instructions,
        "input": user_input,
        "text": {"format": {"type": "text"}, "verbosity": TEXT_VERBOSITY},
        "reasoning": {"effort": REASONING_EFFORT, "summary": REASONING_SUMMARY},
        "store": STORE_RESPONSES,
    }


def call_openai(*, instructions: str, user_input: str, placeholder=None) -> str:
    client = get_openai_client()
    kwargs = build_response_kwargs(instructions, user_input)

    if USE_STREAMING:
        try:
            collected = ""
            events = client.responses.create(**kwargs, stream=True)
            for event in events:
                delta = extract_stream_delta(event)
                if delta:
                    collected += delta
                    if placeholder is not None:
                        placeholder.markdown(collected)
            if collected.strip():
                return collected.strip()
        except Exception as exc:
            if placeholder is not None:
                placeholder.warning(f"스트리밍 응답에 실패하여 일반 응답으로 재시도합니다. ({type(exc).__name__})")

    response = client.responses.create(**kwargs)
    text = response_to_text(response)
    if not text:
        raise RuntimeError("모델 응답이 비어 있습니다. 다시 시도하세요.")
    return text


def generate_initial_summary(mode: str, source_text: str, placeholder) -> str:
    instructions, _ = build_instructions(mode)
    user_input = f"""
다음 회의록 원문을 선택된 프롬프트와 공통 문체 보정 규칙에 따라 요약하십시오.

[회의록 원문]
{source_text}
""".strip()
    return call_openai(instructions=instructions, user_input=user_input, placeholder=placeholder)


def format_chat_history(chat_history: List[Dict[str, str]], limit: int = 8) -> str:
    if not chat_history:
        return "이전 수정 대화 없음"
    rows = chat_history[-limit:]
    return "\n\n".join(f"[{m.get('role', 'unknown')}]\n{m.get('content', '')}" for m in rows)


def needs_source_text_for_revision(user_request: str) -> bool:
    request = user_request.lower()
    keywords = ["원문", "누락", "빠진", "발화자", "참조표", "근거", "표현을 더 살려", "원문 표현", "다시 확인", "검토", "내용 추가", "발언"]
    return any(keyword in request for keyword in keywords)


def revise_summary(mode: str, source_text: str, current_summary: str, chat_history: List[Dict[str, str]], user_request: str, placeholder) -> str:
    extra = """
당신은 기존 요약본을 사용자의 수정 요청에 따라 다듬는 편집자입니다.
반드시 기존 요약본의 구조와 원문 근거를 유지하십시오.
사용자가 요청한 부분만 우선 수정하되, 문체 일관성이 필요하면 전체를 자연스럽게 정리하십시오.
원문에 없는 내용을 새로 만들지 마십시오.
기존 프롬프트가 참조표를 요구하면 참조표를 유지하십시오.
사용자가 “본문만 수정”, “참조표 유지”처럼 범위를 지정하면 그 범위를 우선 따르십시오.
""".strip()
    instructions, _ = build_instructions(mode, extra=extra)

    if needs_source_text_for_revision(user_request):
        user_input = f"""
[원문 회의록]
{source_text}

[현재 요약본]
{current_summary}

[이전 수정 대화]
{format_chat_history(chat_history)}

[이번 사용자 수정 요청]
{user_request}

위 내용을 바탕으로 수정된 최종 요약본을 작성하십시오.
""".strip()
    else:
        user_input = f"""
[현재 요약본]
{current_summary}

[이전 수정 대화]
{format_chat_history(chat_history)}

[이번 사용자 수정 요청]
{user_request}

원문 전체 재검토가 필요한 요청이 아닙니다.
현재 요약본의 구조와 의미를 유지하면서 문체와 표현만 자연스럽게 수정하십시오.
원문에 없는 내용을 새로 만들지 마십시오.
참조표가 있으면 유지하십시오.
""".strip()
    return call_openai(instructions=instructions, user_input=user_input, placeholder=placeholder)


def make_download_file_name(mode: str, ext: str) -> str:
    slug = MODE_SLUG.get(mode, "nanum")
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    return f"{slug}_summary_{ts}.{ext}"


def make_docx_download(text: str) -> bytes:
    doc = Document()
    doc.add_heading("Nanum Meeting Summary", level=1)
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith("##### "):
            doc.add_paragraph(line[6:].strip())
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.getvalue()


def render_sidebar() -> str:
    st.sidebar.title("나눔 회의록 요약 챗봇")
    st.sidebar.caption(f"세션 ID: {st.session_state.session_id[:8]}…")
    st.sidebar.success("인증됨" if st.session_state.get("authenticated") else "인증 필요")

    mode_options = ["나눔 1", "나눔 2"]
    current_mode = st.session_state.get("selected_mode", "나눔 1")
    default_index = mode_options.index(current_mode) if current_mode in mode_options else 0

    # key를 지정하지 않아 세션 초기화 시 widget key 충돌을 방지합니다.
    chosen_mode = st.sidebar.selectbox("요약 유형", options=mode_options, index=default_index)
    st.session_state["selected_mode"] = chosen_mode

    st.sidebar.info(f"모델: {get_model_name()}")

    try:
        prompt_text, prompt_name, version = load_prompt(chosen_mode)
        st.sidebar.caption(f"프롬프트: {prompt_name}")
        st.sidebar.caption(f"버전: {version}")
        with st.sidebar.expander("선택된 프롬프트 미리보기"):
            st.text_area("프롬프트 내용", value=prompt_text, height=260, disabled=True)
    except Exception as exc:
        st.sidebar.error(str(exc))

    if st.sidebar.button("현재 세션 초기화", type="secondary", use_container_width=True):
        reset_current_session(keep_auth=True)
        st.rerun()

    return chosen_mode


def render_input_area() -> str:
    st.subheader("1. 회의록 원문 붙여넣기")
    st.caption("파일 업로드 없이 회의록 원문을 아래 입력창에 붙여넣어 처리합니다.")

    input_key = f"raw_input_area_{st.session_state.get('input_version', 0)}"
    raw_text = st.text_area("회의록 원문", height=360, placeholder="여기에 회의록 원문을 붙여넣으십시오.", key=input_key)
    st.session_state["raw_input_text"] = raw_text
    st.session_state["combined_source_text"] = raw_text.strip()

    if raw_text.strip():
        st.caption(f"처리 대상 텍스트 길이: {len(raw_text.strip()):,}자")
    return raw_text.strip()


def render_summary_area(mode: str) -> None:
    st.subheader("2. 요약 생성")
    source_text = st.session_state.get("combined_source_text", "")

    if st.button("요약 생성", type="primary", use_container_width=True):
        if not source_text or len(source_text.strip()) < 20:
            st.warning("요약할 회의록 원문이 너무 짧거나 비어 있습니다.")
            return
        try:
            _, _, version = load_prompt(mode)
            st.session_state["prompt_version"] = version
        except Exception as exc:
            st.error(f"프롬프트 로딩 오류: {exc}")
            return

        placeholder = st.empty()
        with st.spinner("GPT가 회의록을 요약하는 중입니다..."):
            try:
                summary = generate_initial_summary(mode=mode, source_text=source_text, placeholder=placeholder)
                st.session_state["current_summary"] = summary
                st.session_state["chat_history"] = [{"role": "assistant", "content": summary, "created_at": datetime.now().isoformat(timespec="seconds")}]
                st.success("요약 생성 완료")
                st.rerun()
            except Exception as exc:
                st.error(f"요약 생성 중 오류가 발생했습니다: {exc}")
                st.info("API 키, 모델 접근 권한, 네트워크 상태, rate limit을 확인한 뒤 다시 시도하세요.")

    if st.session_state.get("current_summary"):
        st.subheader("3. 요약 결과")
        st.markdown(st.session_state["current_summary"])


def render_chat_area(mode: str) -> None:
    if not st.session_state.get("current_summary"):
        return

    st.subheader("4. 수정 요청 채팅")
    st.caption("현재 요약본을 바탕으로 수정 요청을 입력하세요. 예: ‘더 자연스럽게’, ‘참조표는 유지하고 본문만 수정’")

    for message in st.session_state.get("chat_history", []):
        role = message.get("role", "assistant")
        if role not in ("user", "assistant"):
            role = "assistant"
        with st.chat_message(role):
            st.markdown(message.get("content", ""))

    user_request = st.chat_input("수정 요청을 입력하세요.")
    if user_request:
        st.session_state["chat_history"].append({"role": "user", "content": user_request, "created_at": datetime.now().isoformat(timespec="seconds")})
        placeholder = st.empty()
        with st.spinner("수정 요청을 반영하는 중입니다..."):
            try:
                revised = revise_summary(
                    mode=mode,
                    source_text=st.session_state.get("combined_source_text", ""),
                    current_summary=st.session_state.get("current_summary", ""),
                    chat_history=st.session_state.get("chat_history", []),
                    user_request=user_request,
                    placeholder=placeholder,
                )
                st.session_state["current_summary"] = revised
                st.session_state["chat_history"].append({"role": "assistant", "content": revised, "created_at": datetime.now().isoformat(timespec="seconds")})
                st.rerun()
            except Exception as exc:
                st.error(f"수정 중 오류가 발생했습니다: {exc}")
                st.info("잠시 후 다시 시도하세요. rate limit 또는 네트워크 문제일 수 있습니다.")


def render_download_area(mode: str) -> None:
    if not st.session_state.get("current_summary"):
        return

    st.subheader("5. 다운로드")
    summary = st.session_state["current_summary"]

    col1, col2, col3 = st.columns(3)
    with col1:
        st.download_button("Markdown 다운로드", data=summary.encode("utf-8"), file_name=make_download_file_name(mode, "md"), mime="text/markdown", use_container_width=True)
    with col2:
        st.download_button("TXT 다운로드", data=summary.encode("utf-8"), file_name=make_download_file_name(mode, "txt"), mime="text/plain", use_container_width=True)
    with col3:
        st.download_button("DOCX 다운로드", data=make_docx_download(summary), file_name=make_download_file_name(mode, "docx"), mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="📝", layout="wide")
    init_session_state()

    if not st.session_state.get("authenticated"):
        check_password()
        return

    mode = render_sidebar()

    st.title("회의록 요약")
    st.write("나눔 1 또는 나눔 2를 선택한 뒤 회의록 원문을 붙여넣으십시오.")
    st.caption("파일 업로드 기능은 제거되어 있으며, 붙여넣은 텍스트만 처리합니다.")

    try:
        _, prompt_name, prompt_version_value = load_prompt(mode)
        st.info(f"현재 요약 유형: {mode} · 프롬프트: {prompt_name} · 버전: {prompt_version_value}")
    except Exception as exc:
        st.error(f"프롬프트 파일 오류: {exc}")
        st.stop()

    render_input_area()
    render_summary_area(mode)
    render_chat_area(mode)
    render_download_area(mode)


if __name__ == "__main__":
    main()
